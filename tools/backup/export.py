#!/usr/bin/env python3
"""
Bind Build ERP — export everything to CSV, Excel and PDF.

Written for a Google Drive backup: point this at a folder that Google Drive
for Desktop syncs, schedule it, and you get an off-site copy in formats you
can open without a database.

This complements tools/backup/backup.sh, it does not replace it:

    backup.sh   pg_dump   -> can be RESTORED into a working system
    export.py   csv/xlsx  -> can be READ by a human, an accountant, Excel

You need both. A CSV of your invoices is useless for recovery; a pg_dump is
useless to your CA.

Usage
-----
    export SUPABASE_DB_URL="postgresql://postgres.[ref]:[pw]@aws-0-ap-south-1.pooler.supabase.com:5432/postgres"
    python3 tools/backup/export.py --out "/Users/you/Google Drive/My Drive/BindBuild Backups"

Options
-------
    --out DIR        destination (default ./exports)
    --skip-empty     omit tables with no rows
    --no-pdf         skip the PDF summary
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import sys
from pathlib import Path

try:
    import psycopg2
    import psycopg2.extras
except ImportError:
    sys.exit("psycopg2 is missing.  pip install psycopg2-binary")

import csv

# Excel caps sheet names at 31 chars and forbids these characters
INVALID_SHEET = re.compile(r'[\[\]\*/\\?:]')

# tables an accountant actually asks for, in the order they ask
PRIORITY = [
    'invoices', 'invoice_items', 'invoice_payments', 'credit_notes',
    'expenses', 'purchase_orders', 'po_items',
    'payroll_runs', 'payslips', 'salary_structures',
    'clients', 'projects', 'leads', 'proposals',
]


def connect():
    url = os.environ.get('SUPABASE_DB_URL')
    if not url:
        sys.exit("SUPABASE_DB_URL is not set. See the docstring at the top of this file.")
    return psycopg2.connect(url)


def list_tables(cur):
    cur.execute("""
        select table_name
        from information_schema.tables
        where table_schema = 'public' and table_type = 'BASE TABLE'
        order by table_name
    """)
    names = [r[0] for r in cur.fetchall()]
    # priority tables first so the workbook opens on something useful
    return [t for t in PRIORITY if t in names] + [t for t in names if t not in PRIORITY]


def fetch(cur, table):
    cur.execute(f'select * from public."{table}"')
    cols = [d[0] for d in cur.description]
    return cols, cur.fetchall()


def sheet_name(table, used):
    name = INVALID_SHEET.sub('_', table)[:31]
    n = 1
    while name.lower() in used:
        suffix = f'_{n}'
        name = name[:31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def write_csvs(cur, tables, out: Path, skip_empty: bool):
    csv_dir = out / 'csv'
    csv_dir.mkdir(parents=True, exist_ok=True)
    counts = {}
    for t in tables:
        cols, rows = fetch(cur, t)
        if skip_empty and not rows:
            continue
        counts[t] = len(rows)
        with open(csv_dir / f'{t}.csv', 'w', newline='', encoding='utf-8-sig') as fh:
            w = csv.writer(fh)
            w.writerow(cols)
            w.writerows(rows)
    return counts


def write_xlsx(cur, tables, out: Path, counts: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used, head_fill = set(), PatternFill('solid', fgColor='1F2933')
    head_font = Font(bold=True, color='FFFFFF')

    # contents sheet first
    toc = wb.create_sheet('Contents')
    toc.append(['Table', 'Rows'])
    for c in toc[1]:
        c.font, c.fill = head_font, head_fill

    for t in tables:
        if t not in counts:
            continue
        cols, rows = fetch(cur, t)
        ws = wb.create_sheet(sheet_name(t, used))
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical='center')
        for r in rows:
            # Excel cannot hold tz-aware datetimes or dicts
            ws.append([
                (v.replace(tzinfo=None) if isinstance(v, dt.datetime) and v.tzinfo else
                 str(v) if isinstance(v, (dict, list)) else v)
                for v in r
            ])
        ws.freeze_panes = 'A2'
        for i, col in enumerate(cols, 1):
            width = max(len(str(col)) + 2, 12)
            ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
        toc.append([t, counts[t]])

    toc.column_dimensions['A'].width = 34
    toc.column_dimensions['B'].width = 10
    path = out / 'bindbuild-data.xlsx'
    wb.save(path)
    return path


def write_pdf(cur, out: Path, stamp: str):
    """A one-page summary an accountant can read without opening anything else."""
    from fpdf import FPDF

    def scalar(sql, default=0):
        try:
            cur.execute(sql)
            v = cur.fetchone()
            return v[0] if v and v[0] is not None else default
        except Exception:
            cur.connection.rollback()
            return default

    billed    = scalar("select sum(total) from public.invoices where deleted_at is null")
    received  = scalar("select sum(amount) from public.invoice_payments where deleted_at is null")
    credited  = scalar("select sum(total) from public.credit_notes where deleted_at is null")
    cost      = scalar("select sum(amount) from public.expenses where status='paid' and deleted_at is null")
    projects  = scalar("select count(*) from public.projects where deleted_at is null")
    clients   = scalar("select count(*) from public.clients where deleted_at is null")
    invoices  = scalar("select count(*) from public.invoices where deleted_at is null")
    employees = scalar("select count(*) from public.employees where deleted_at is null")
    payslips  = scalar("select count(*) from public.payslips where deleted_at is null")

    def inr(v):
        return 'Rs ' + f'{float(v or 0):,.2f}'

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font('Helvetica', 'B', 18)
    pdf.cell(0, 12, 'Bind Build ERP - data export', new_x='LMARGIN', new_y='NEXT')
    pdf.set_font('Helvetica', '', 10)
    pdf.set_text_color(110, 110, 110)
    pdf.cell(0, 6, f'Studio Bind Architects  |  generated {stamp}',
             new_x='LMARGIN', new_y='NEXT')
    pdf.ln(6)

    def section(title, rows):
        pdf.set_text_color(20, 20, 20)
        pdf.set_font('Helvetica', 'B', 12)
        pdf.cell(0, 9, title, new_x='LMARGIN', new_y='NEXT')
        pdf.set_font('Helvetica', '', 10)
        for label, value in rows:
            pdf.set_text_color(110, 110, 110)
            pdf.cell(70, 7, label)
            pdf.set_text_color(20, 20, 20)
            pdf.cell(0, 7, str(value), new_x='LMARGIN', new_y='NEXT')
        pdf.ln(4)

    section('Money', [
        ('Invoiced',              inr(billed)),
        ('Received',              inr(received)),
        ('Credit notes raised',   inr(credited)),
        ('Outstanding',           inr(float(billed or 0) - float(received or 0) - float(credited or 0))),
        ('Cost booked (paid)',    inr(cost)),
    ])
    section('Records', [
        ('Clients',   clients),
        ('Projects',  projects),
        ('Invoices',  invoices),
        ('Employees', employees),
        ('Payslips',  payslips),
    ])

    pdf.set_font('Helvetica', 'I', 9)
    pdf.set_text_color(130, 130, 130)
    pdf.multi_cell(0, 5,
        'This is a readable summary, not a restorable backup. Use '
        'tools/backup/backup.sh for anything you would need to restore from. '
        'Figures exclude soft-deleted records.')

    path = out / 'summary.pdf'
    pdf.output(str(path))
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', default='./exports')
    ap.add_argument('--skip-empty', action='store_true')
    ap.add_argument('--no-pdf', action='store_true')
    args = ap.parse_args()

    stamp = dt.datetime.now().strftime('%Y-%m-%d_%H%M')
    out = Path(args.out).expanduser() / stamp
    out.mkdir(parents=True, exist_ok=True)

    conn = connect()
    cur = conn.cursor()

    tables = list_tables(cur)
    print(f'{len(tables)} tables')

    counts = write_csvs(cur, tables, out, args.skip_empty)
    print(f'  csv     {len(counts)} files, {sum(counts.values())} rows')

    xlsx = write_xlsx(cur, tables, out, counts)
    print(f'  xlsx    {xlsx.name}')

    if not args.no_pdf:
        pdf = write_pdf(cur, out, stamp)
        print(f'  pdf     {pdf.name}')

    cur.close()
    conn.close()

    print(f'\nWritten to {out}')
    print('If this folder is inside Google Drive, it syncs on its own.')


if __name__ == '__main__':
    main()
